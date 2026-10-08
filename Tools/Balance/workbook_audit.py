#!/usr/bin/env python3
"""Read-only workbook inventory used by the balance source audit."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def cell_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    output = []
    for path in sorted((root / "Assets").rglob("*.xlsx")):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        entry = {"path": path.relative_to(root).as_posix(), "sheets": []}
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [cell_value(value) for value in row]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            entry["sheets"].append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "formula_count": sum(
                    1 for row in rows for value in row if isinstance(value, str) and value.startswith("=")
                ),
                "rows": rows,
            })
        output.append(entry)
    text = json.dumps(output, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
